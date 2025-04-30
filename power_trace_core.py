import pandas as pd
import sys
import os
import re
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

# 參數：EXP_Processed.xlsx, 深度, power_nets字串, summary_path, trace_path
input_file = sys.argv[1]
max_depth = int(sys.argv[2])
source_input = sys.argv[3]
summary_output_file = sys.argv[4]
tracking_output_file = sys.argv[5]

source_nets = [x.strip() for x in source_input.split(',') if x.strip()]

# ========== 資料處理 ==========
df = pd.read_excel(input_file, engine="openpyxl")
df.fillna('', inplace=True)

for col in ['ID', 'Part Reference', 'Value', 'Net Name']:
    df[col] = df[col].astype(str).str.strip().str.upper()

pin_info = {}
part_to_pins = {}

for _, row in df.iterrows():
    part = row['Part Reference']
    id_value = row['ID']
    netname = row['Net Name']
    if ':' in id_value:
        parts = id_value.split(':')
        if len(parts) == 2:
            pin_number, pin_name = parts[1].split(' ', 1) if ' ' in parts[1] else (parts[1], parts[1])
            pin_info[(part, pin_number.strip())] = (pin_name.strip(), netname.strip())
            part_to_pins.setdefault(part, []).append(pin_number.strip())

initial_power_candidates = [
    net for net in df['Net Name'].unique()
    if isinstance(net, str) and net.startswith('+') and ('V' in net or 'VA' in net)
]
blacklist_keywords = ['_PWRGD','_PG','_PGOOD','_POK','_PWROK','_EN','_FB','_REF','_VOS','_CLM','LIM','RS','RSO','ROS','_BST','_BOOT','_BOOST','_MODE','_SW','_CS','_VSNS','_PHASE','_VDRV','_RGND','_TRK','_RUN','_VFB','_ITH','_TG','_BG','_SENS','_RAMP']
true_power_nets = [net for net in initial_power_candidates if not any(k in net for k in blacklist_keywords)]

def is_ground_net(netname):
    return any(netname.upper().startswith(g) for g in ['GND', 'PGND', 'AGND', 'SGND'])

def is_ic_power_output_pin(pin_name):
    allowed_keywords = ['VTT', 'VOUT', 'REGOUT', 'SW', 'DRV', 'OUT','OUT_', 'LX', 'INTVCC', 'CTRL']
    pin_upper = pin_name.upper()
    for keyword in allowed_keywords:
        idx = pin_upper.find(keyword)
        if idx != -1:
            suffix = pin_upper[idx + len(keyword):]
            if suffix == '' or suffix[0].isdigit():
                return True
    return False

def is_standard_part(partref, prefixes):
    if (partref.startswith('S') or partref.startswith('P')) and partref[1:2].isdigit():
        return True
    return any(partref.startswith(prefix) and partref[len(prefix):][:1].isdigit() for prefix in prefixes)

summary_records = []
trace_detail_dict = {}

def trace_single_net(start_net):
    trace_records = []
    power_source_found = False
    visited_pins = set()
    visited_nets = set()
    visited_paths = set()
    start_pins = [(part, pin_num, pin_name) for (part, pin_num), (pin_name, netname) in pin_info.items() if netname == start_net]

    def trace_from_pin(partref, pin_number, path, visited_pins, visited_nets, visited_paths, depth):
        nonlocal power_source_found
        if (partref, pin_number) not in pin_info:
            return
        pin_name, netname = pin_info[(partref, pin_number)]
        path_id = (netname, partref, pin_number)
        if path_id in visited_paths or depth > max_depth:
            return
        visited_paths.add(path_id)
        visited_pins.add((partref, pin_number))
        visited_nets.add(netname)
        if is_ground_net(netname):
            path.append(("Net", netname, "接到GND(停止)"))
            trace_records.append(path)
            return
        if netname in source_nets or (netname in true_power_nets and netname != start_net):
            path.append(("Net", netname, f"成功連到Power Net {netname}！（停止）"))
            trace_records.append(path)
            power_source_found = True
            return
        if not path or path[-1][1] != netname:
            path.append(("Net", netname, ""))
        for (other_partref, other_pinnum), (other_pinname, other_netname) in pin_info.items():
            if other_netname != netname or (other_partref, other_pinnum) in visited_pins:
                continue
            if other_partref.startswith(('C', 'SC')):
                trace_records.append(path + [("Part", f"{other_partref}({other_pinnum})", "遇到電容(停止)")])
                continue
            if other_partref.startswith(('U', 'SU')):
                if is_ic_power_output_pin(other_pinname):
                    trace_records.append(path + [("Part", f"{other_partref}({other_pinnum})", f"IC Power Output Pin ({other_pinname}) 找到Power Source停止")])
                    power_source_found = True
                    return
                else:
                    trace_records.append(path + [("Part", f"{other_partref}({other_pinnum})", "遇到IC(非Power腳)(停止)")])
                    continue
            if is_standard_part(other_partref, ['L','SL','R','SR','D','SD','F','SF','FB','SFB','Q','SQ','JP','SJP']):
                for next_pinnum in part_to_pins.get(other_partref, []):
                    if (other_partref, next_pinnum) != (other_partref, other_pinnum):
                        next_net = pin_info.get((other_partref, next_pinnum), ('', ''))[1]
                        if next_net and next_net != netname:
                            path_cross = path + [
                                ("Part", f"{other_partref}({other_pinnum})", "跨到"),
                                ("Part", f"{other_partref}({next_pinnum})", "跨到"),
                                ("Net", next_net, "")
                            ]
                            trace_from_pin(other_partref, next_pinnum, path_cross, visited_pins.copy(), visited_nets.copy(), visited_paths.copy(), depth + 1)
            else:
                trace_records.append(path + [("Part", f"{other_partref}({other_pinnum})", "非標準元件(停止)")])

    for part, pin_num, pin_name in start_pins:
        if part.startswith(('U', 'SU')) and is_ic_power_output_pin(pin_name):
            trace_records.append([("Net", start_net, ""), ("Part", f"{part}({pin_num})", f"IC Power Output Pin ({pin_name}) 找到Power Source停止")])
            power_source_found = True
            break
        trace_from_pin(part, pin_num, [], set(), set(), set(), 1)
        if power_source_found:
            break
    if not power_source_found:
        summary_records.append({"Power Net": start_net, "Result": "No Power Source"})
    trace_detail_dict[start_net] = trace_records

for net in true_power_nets:
    trace_single_net(net)

summary_df = pd.DataFrame(summary_records)
wb = Workbook()
ws = wb.active
ws.title = "Summary"
ws.append(["Power Net", "Result"])
for r in summary_df.itertuples(index=False):
    ws.append(r)
for col in ws.columns:
    max_len = max(len(str(cell.value)) if cell.value else 0 for cell in col)
    ws.column_dimensions[get_column_letter(col[0].column)].width = max_len + 5
    for cell in col:
        cell.font = Font(name='Calibri', size=12)
wb.save(summary_output_file)

wb_tracking = Workbook()
ws0 = wb_tracking.active
ws0.title = "Summary"
for net, records in trace_detail_dict.items():
    if not records:
        continue
    ws = wb_tracking.create_sheet(title=re.sub(r'[\\/*?:\[\]]', '_', net)[:31])
    ws.append(["Type", "Name", "Note"])
    for step in records:
        for t, n, note in step:
            ws.append([t, n, note])
wb_tracking.save(tracking_output_file)
