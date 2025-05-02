from flask import Flask, render_template, request, send_file
import pandas as pd
import os
import re
from datetime import datetime
from openpyxl import load_workbook
from openpyxl.styles import Font

app = Flask(__name__)
UPLOAD_FOLDER = "uploads"
RESULT_FOLDER = "results"
os.makedirs(UPLOAD_FOLDER, exist_ok=True)
os.makedirs(RESULT_FOLDER, exist_ok=True)

def process_exp_file(file_path):
    with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
        lines = [line.strip() for line in f if line.strip()]
    headers = [h.strip('"') for h in lines[1].split('\t')]
    data_lines = lines[2:]
    data = [[cell.strip('"') for cell in row.split('\t')] for row in data_lines]
    max_cols = len(headers)
    data = [row[:max_cols] + [''] * (max_cols - len(row)) for row in data]
    df = pd.DataFrame(data, columns=headers)

    def clean_illegal_chars(val):
        if isinstance(val, str):
            return re.sub(r'[\x00-\x08\x0B-\x0C\x0E-\x1F\x7F]', '', val)
        return val

    df = df.applymap(clean_illegal_chars)
    wanted_columns = ["ID", "Part Reference", "Value", "Net Name"]
    def normalize(col_name):
        return col_name.strip().lower().replace(" ", "").replace("_", "")

    normalized_map = {normalize(col): col for col in df.columns}
    final_cols = [normalized_map[nc] for nc in map(normalize, wanted_columns) if nc in normalized_map]
    df = df[final_cols]

    df['Part Reference'] = df['Part Reference'].replace({'<null>': '', None: ''})
    df['Value'] = df['Value'].replace({'<null>': '', None: ''})

    main_part_ref_map, main_value_map = {}, {}
    new_rows = []
    for _, row in df.iterrows():
        id_value = str(row['ID']).strip()
        if ':' not in id_value:
            main_part_ref_map[id_value] = row['Part Reference']
            main_value_map[id_value] = row['Value']
        else:
            main_id = id_value.split(':')[0]
            if row['Part Reference'] == '':
                row['Part Reference'] = main_part_ref_map.get(main_id, '')
            row['Value'] = main_value_map.get(main_id, row['Value'])
            new_rows.append(row)
    return pd.DataFrame(new_rows)

def compare_df(df1, df2, primary_key, ignore_case_space):
    if ignore_case_space:
        df1 = df1.applymap(lambda x: str(x).strip().upper())
        df2 = df2.applymap(lambda x: str(x).strip().upper())
    df1.set_index(primary_key, inplace=True)
    df2.set_index(primary_key, inplace=True)
    df1 = df1[~df1.index.duplicated(keep='first')]
    df2 = df2[~df2.index.duplicated(keep='first')]

    added_keys = df2.index.difference(df1.index)
    removed_keys = df1.index.difference(df2.index)
    common_keys = df1.index.intersection(df2.index)
    df_added = df2.loc[added_keys].copy()
    df_removed = df1.loc[removed_keys].copy()
    df_modified = []
    for key in common_keys:
        old_row = df1.loc[key]
        new_row = df2.loc[key]
        diff = old_row != new_row
        if diff.any():
            combined = new_row.copy()
            for col in df1.columns:
                if diff[col]:
                    combined[col] = f"{old_row[col]} → {new_row[col]}"
            combined[primary_key] = key
            df_modified.append(combined)
    df_modified = pd.DataFrame(df_modified)
    return df_added, df_removed, df_modified

def autofit_columns(ws):
    for column_cells in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in column_cells)
        col = column_cells[0].column_letter
        ws.column_dimensions[col].width = max_len + 2

def export_to_excel(df_added, df_removed, df_modified):
    timestamp = datetime.now().strftime("%Y%m%d_%H%M")
    filename = os.path.join(RESULT_FOLDER, f"EXP_Compare_Report_{timestamp}.xlsx")
    with pd.ExcelWriter(filename, engine="openpyxl") as writer:
        df_added.reset_index().to_excel(writer, sheet_name="新增列", index=False)
        df_removed.reset_index().to_excel(writer, sheet_name="刪除列", index=False)
        df_modified = df_modified.reset_index().rename(columns={"index": "ID"})
        df_modified.to_excel(writer, sheet_name="修改列", index=False)
    wb = load_workbook(filename)
    for sheet in ["新增列", "刪除列", "修改列"]:
        if sheet in wb.sheetnames:
            ws = wb[sheet]
            autofit_columns(ws)
            ws.freeze_panes = "A2"
            if sheet == "修改列":
                red_font = Font(color="FF0000")
                for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=ws.max_column):
                    for cell in row:
                        if isinstance(cell.value, str) and '→' in cell.value:
                            cell.font = red_font
    wb.save(filename)
    return filename

@app.route('/', methods=['GET', 'POST'])
def index():
    if request.method == 'POST':
        file1 = request.files['exp1']
        file2 = request.files['exp2']
        key = request.form['primary_key']
        ignore_case = request.form.get('ignore_case') == 'on'
        path1 = os.path.join(UPLOAD_FOLDER, file1.filename)
        path2 = os.path.join(UPLOAD_FOLDER, file2.filename)
        file1.save(path1)
        file2.save(path2)
        df1 = process_exp_file(path1)
        df2 = process_exp_file(path2)
        df_added, df_removed, df_modified = compare_df(df1, df2, key, ignore_case)
        result_path = export_to_excel(df_added, df_removed, df_modified)
        return send_file(result_path, as_attachment=True)
    return render_template('index.html')

if __name__ == '__main__':
    app.run(debug=True, host="0.0.0.0", port=5000)
